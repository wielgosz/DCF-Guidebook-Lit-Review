"""Tests for the v2.3 desktop runner engine.

Run from desktop_runner/runner:  python -m pytest tests -q
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HERE))

from protocol_engine import __version__  # noqa: E402
from protocol_engine.corpus_register import parse_apa, title_match_score  # noqa: E402
from protocol_engine.figures import render_term_figures  # noqa: E402
from protocol_engine.publication_pack import build_publication_pack  # noqa: E402
from protocol_engine.reference_tables import (  # noqa: E402
    packaged_tables_dir, read_reference_tables, validate_reference_tables,
)
from protocol_engine.run_protocol import PACKAGED_STYLE, PACKAGED_TEMPLATE  # noqa: E402
from protocol_engine.tables_version import describe, read_version, tables_version_for  # noqa: E402


def test_packaged_tables_version_matches_tool_version():
    stamp = read_version(packaged_tables_dir())
    assert stamp["tool_version"] == __version__
    assert stamp["tables_version"] == tables_version_for(__version__)


def test_packaged_tables_unchanged_since_stamp():
    info = describe(packaged_tables_dir())
    assert info["state"] == "unchanged", (
        f"Packaged tables edited without a release ({info['changed_files']}); "
        "run tools/release_tables.py --notes '...'"
    )


def test_tables_version_scheme():
    assert tables_version_for("2.3.0") == "2.30"
    assert tables_version_for("2.3.1") == "2.31"
    assert tables_version_for("2.4.0") == "2.40"


def test_packaged_tables_have_no_structural_errors():
    folder = packaged_tables_dir()
    data = read_reference_tables(folder)
    errors = [i for i in validate_reference_tables(folder, data, need_datasets=True) if i.severity == "ERROR"]
    assert not errors, errors


def test_edited_copy_is_labelled_custom(tmp_path):
    copy = tmp_path / "tables"
    shutil.copytree(packaged_tables_dir(), copy)
    with (copy / "exclusions_duplicates.csv").open("a", encoding="utf-8-sig") as f:
        f.write("extra.pdf,EXCLUDE_DUPLICATE,,test\n")
    info = describe(copy)
    assert info["state"] == "modified" and "exclusions_duplicates.csv" in info["changed_files"]


@pytest.mark.parametrize("apa, year, title, url", [
    ("ABIOVE. (2024). Cerrado Monitoring Report. https://abiove.org.br/x/", "2024", "Cerrado Monitoring Report", "https://abiove.org.br/x/"),
    ("Trase. (n.d.). Supply chain mapping manual. https://trase.earth.", "", "Supply chain mapping manual", "https://trase.earth"),
    ("WWF. (accessed 2026). DCF Toolkit. https://wwf.org/a", "2026", "DCF Toolkit", "https://wwf.org/a"),
    ("Deutsche Gesellschaft für Internationale Zusammenarbeit (GIZ) GmbH, & Olab. (2026). AB+S dry-run. https://e.eu/x.pdf",
     "2026", "AB+S dry-run", "https://e.eu/x.pdf"),
])
def test_parse_apa(apa, year, title, url):
    p = parse_apa(apa)
    assert (p["year"], p["title"], p["url"]) == (year, title, url)


def test_title_match_score():
    assert title_match_score("Cerrado Monitoring Report", "ABIOVE cerrado monitoring report 2023") == 1.0
    assert title_match_score("Common Guidance for the Identification of HCV", "Operational guidance monitoring") < 0.5


def _term_summary():
    return pd.DataFrame({
        "category": ["Jurisdictional terms", "Jurisdictional terms", "Supply chain terms", "Farm level terms"],
        "term": ["landscape", "jurisdiction", "trader", "farm"],
        "reports_referencing": [10, 4, 7, 9],
    })


def test_figures_are_byte_reproducible(tmp_path):
    a = render_term_figures(_term_summary(), tmp_path / "a", PACKAGED_STYLE)
    b = render_term_figures(_term_summary(), tmp_path / "b", PACKAGED_STYLE)
    assert len(a) == 6 and a == b


def test_publication_pack_preserves_rdi_header_rows(tmp_path):
    d1 = pd.DataFrame({"category": ["Jurisdictional terms"], "term": ["landscape"], "variants_included": ["landscape"],
                       "reports_referencing": [10], "total_occurrences": [42], "rank_in_category": [1]})
    b1 = pd.DataFrame({"apa_reference": ["Org. (2024). Title. https://e.org"], "source_url": ["https://e.org"]})
    a1 = pd.DataFrame({"organization": ["Org"], "website_url": ["https://e.org"]})
    out = tmp_path / "pack.xlsx"
    render_term_figures(_term_summary(), tmp_path / "fig", PACKAGED_STYLE)
    res = build_publication_pack(PACKAGED_TEMPLATE, out, d1=d1, a1=a1, b1=b1, figure_dir=tmp_path / "fig", run_info={})
    assert any(r.startswith("Table D1") for r in res["refreshed"])
    src, dst = load_workbook(PACKAGED_TEMPLATE), load_workbook(out)
    for tab in ["Figure 1", "Table A1", "Table B1", "Table D1"]:
        for row in range(1, 10):
            for col in range(1, 4):
                assert src[tab].cell(row, col).value == dst[tab].cell(row, col).value, (tab, row, col)
    values = [c.value for row in dst["Table D1"].iter_rows() for c in row]
    assert "landscape" in values and 42 in values
    assert "Run Info" in dst.sheetnames


def test_selftest_end_to_end():
    from protocol_engine.selftest import run_selftest
    assert run_selftest() == 0
