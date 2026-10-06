"""Build the generic publication template from an RDI-format figure template.

    python tools/build_publication_template.py --from "FIGURE-TEMPLATE ... (v8.0).xlsx"

Writes templates/Publication_Template.xlsx: a copy of the source workbook that
keeps, for maximum compatibility with the RDI copy-edit format,

- the READ ME tab;
- each standard tab's header block (caption, y/x-axis units, Source, Source url,
  RDI/Comms/CopyEdit Requirements) exactly as in the source;
- the "INSERT FIGURE HERE" anchor and each table's header row, in place;

and removes everything else: all rows below each table header (the data),
images, and project-specific tabs. The runner fills a copy of this template
(or of the user's own RDI template) at run time.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HERE))
from protocol_engine.publication_pack import (  # noqa: E402
    A1_SPEC, B1_HEADERS, D1_SPEC, FIGURE_ANCHOR, find_cell, find_header,
)

OUT = HERE / "templates" / "Publication_Template.xlsx"
KEEP_TABS = ["READ ME", "Figure 1", "Figure 2", "Figure 3",
             "Table A1", "Table B1", "Table C1", "Table D1", "Table E1"]
C1_HEADERS = [["dcf relevance"], ["dataset name"], ["apa citation"]]
E1_HEADERS = [["figure"], ["dataset name"], ["coverage level"]]


def last_layout_row(ws) -> int:
    """Row of the last layout element (anchor, heading or table header)."""
    rows = []
    for labels in ([o for o, _ in D1_SPEC.columns], [o for o, _ in A1_SPEC.columns], C1_HEADERS, E1_HEADERS):
        found = find_header(ws, labels)
        if found:
            rows.append(found[0])
    for label in B1_HEADERS + [FIGURE_ANCHOR]:
        cell = find_cell(ws, label)
        if cell is not None:
            rows.append(cell.row)
    return max(rows) if rows else 9


def strip_data(ws) -> int:
    keep_to = last_layout_row(ws)
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row > keep_to:
            ws.unmerge_cells(str(rng))
    if ws.max_row > keep_to:
        ws.delete_rows(keep_to + 1, ws.max_row - keep_to)
    ws._images = []
    ws._charts = []
    return keep_to


def main() -> Path:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--from", dest="source", required=True, help="RDI-format figure template (.xlsx) to derive from.")
    ap.add_argument("--out", default=str(OUT))
    args = ap.parse_args()

    wb = load_workbook(args.source)
    missing = [t for t in KEEP_TABS if t not in wb.sheetnames]
    if missing:
        raise SystemExit(f"Source template lacks tab(s): {', '.join(missing)}")
    for name in list(wb.sheetnames):
        if name not in KEEP_TABS:
            del wb[name]
    for name in KEEP_TABS:
        if name == "READ ME":
            continue
        kept = strip_data(wb[name])
        print(f"{name}: kept rows 1-{kept}")
    wb._sheets = [wb[n] for n in KEEP_TABS]
    wb.active = 0
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(main())
