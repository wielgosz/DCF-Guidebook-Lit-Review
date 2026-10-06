"""Corpus register (Table B1) validation.

The B1_Corpus_Documents tab of the input workbook is the corpus register: it
links every PDF file name to a stable doc_id and its APA reference. Updating it
is the most common maintenance step, so it can be validated on its own
(``validate-register``) before a full protocol run.

Checks per register row
-----------------------
- file_name present, unique, and found in the PDF folder (or ZIPs within it)
- not listed as excluded/duplicate in Exclusions_Duplicates
- doc_id unique; publishing_org_id exists in A1
- APA reference present; its year matches ``year``; it contains ``title``;
  its URL matches ``source_url``; ``source_url`` is a well-formed http(s) URL
- the title actually appears on the first pages of the PDF (catches rows whose
  file and metadata have drifted apart)
- identical PDF content registered twice
- optional: every URL resolves (``check_links``)
- optional: the APA reference appears in the publication template's Table B1

PDFs in the folder that are not in the register (and not excluded) are listed
with metadata guesses so they can be added. They are never counted.
"""
from __future__ import annotations

import hashlib
import re
import threading
import unicodedata
import urllib.error
import urllib.request
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict, dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

Log = Callable[[str], None]

APA_RE = re.compile(r"^(?P<author>.+?)\s*\((?P<date>[^)]*)\)\.?\s*(?P<rest>.*)$", re.S)
URL_RE = re.compile(r"https?://[^\s<>\"]+", re.I)
YEAR_RE = re.compile(r"(19|20)\d{2}")
STOPWORDS = {
    "with", "from", "into", "that", "this", "their", "about", "under", "guide",
    "the", "and", "for", "les", "des", "para", "como", "uma", "dos", "das",
}
PROBE_PAGES = 3
PROBE_TIMEOUT_SECONDS = 90
TITLE_MATCH_MIN = 0.5

STATUS_ORDER = {"ERROR": 0, "CHECK": 1, "OK": 2, "REFERENCE_ONLY": 3}
STATUS_FILL = {
    "ERROR": PatternFill("solid", fgColor="F8D7DA"),
    "CHECK": PatternFill("solid", fgColor="FFF3CD"),
    "OK": PatternFill("solid", fgColor="D4EDDA"),
    "REFERENCE_ONLY": PatternFill("solid", fgColor="E2E3E5"),
}


@dataclass
class RegisterIssue:
    severity: str  # ERROR | WARNING | INFO
    doc_id: str
    file_name: str
    check: str
    message: str


@dataclass
class PdfProbe:
    path: Path
    sha256: str = ""
    pages: Optional[int] = None
    first_text: str = ""
    meta_title: str = ""
    error: str = ""


@dataclass
class RegisterResult:
    register: pd.DataFrame
    manifest: pd.DataFrame
    unregistered: pd.DataFrame
    issues: List[RegisterIssue]
    publication_recon: pd.DataFrame = field(default_factory=pd.DataFrame)

    @property
    def errors(self) -> List[RegisterIssue]:
        return [i for i in self.issues if i.severity == "ERROR"]

    @property
    def warnings(self) -> List[RegisterIssue]:
        return [i for i in self.issues if i.severity == "WARNING"]


# --------------------------------------------------------------------------
# Text helpers
# --------------------------------------------------------------------------

def fold(text: str) -> str:
    """Lower-case, strip accents and punctuation, collapse whitespace."""
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^\w\s]", " ", text.casefold())
    return re.sub(r"\s+", " ", text).strip()


def normalize_apa(text: str) -> str:
    text = unicodedata.normalize("NFKC", str(text or "")).casefold()
    return re.sub(r"\s+", " ", text).strip().rstrip(".")


def normalize_url(url: str) -> str:
    url = str(url or "").strip().rstrip(".,;)")
    url = re.sub(r"^https?://(www\.)?", "", url, flags=re.I)
    return url.rstrip("/").lower()


def parse_apa(apa: str) -> Dict[str, str]:
    apa = str(apa or "").strip()
    m = APA_RE.match(apa)
    urls = URL_RE.findall(apa)
    url = urls[-1].rstrip(".,;)") if urls else ""
    if not m:
        return {"author": "", "date": "", "year": "", "title": "", "url": url}
    rest = m.group("rest")
    if url:
        rest = rest.split(urls[-1])[0]
    year = YEAR_RE.search(m.group("date"))
    return {
        "author": m.group("author").strip().rstrip("."),
        "date": m.group("date").strip(),
        "year": year.group(0) if year else "",
        "title": rest.strip().rstrip("."),
        "url": url,
    }


def clean_year(value: str) -> str:
    m = YEAR_RE.search(str(value or ""))
    return m.group(0) if m else ""


def title_match_score(title: str, haystack: str) -> Optional[float]:
    """Share of significant title words found in ``haystack`` (None if no words)."""
    words = {w for w in fold(title).split() if len(w) >= 4 and w not in STOPWORDS}
    if not words:
        return None
    hay = set(fold(haystack).split())
    return round(sum(1 for w in words if w in hay) / len(words), 2)


# --------------------------------------------------------------------------
# PDF discovery and probing
# --------------------------------------------------------------------------

def discover_pdfs(pdf_folder: Path, unzip_dir: Path) -> Dict[str, List[Path]]:
    """Map lower-cased file name -> PDF paths (loose files plus ZIP contents)."""
    found: Dict[str, List[Path]] = {}
    for p in sorted(pdf_folder.rglob("*")):
        if not p.is_file():
            continue
        if p.suffix.lower() == ".zip":
            target = unzip_dir / p.stem.replace(" ", "_")
            target.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(p) as z:
                z.extractall(target)
            for q in sorted(target.rglob("*.pdf")):
                found.setdefault(q.name.lower(), []).append(q)
        elif p.suffix.lower() == ".pdf":
            found.setdefault(p.name.lower(), []).append(p)
    return found


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def probe_pdf(path: Path) -> PdfProbe:
    """Hash the file and read its first pages + metadata title (bounded time)."""
    probe = PdfProbe(path=path, sha256=sha256_file(path))
    box: Dict[str, object] = {}

    def work() -> None:
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            box["pages"] = len(reader.pages)
            texts = []
            for page in reader.pages[:PROBE_PAGES]:
                texts.append(page.extract_text() or "")
            box["text"] = "\n".join(texts)
            meta = reader.metadata or {}
            box["title"] = str(meta.get("/Title", "") or "")
        except Exception as exc:  # pragma: no cover - depends on the PDF
            box["error"] = f"{type(exc).__name__}: {exc}"

    t = threading.Thread(target=work, daemon=True)
    t.start()
    t.join(PROBE_TIMEOUT_SECONDS)
    if t.is_alive():
        probe.error = f"timed out after {PROBE_TIMEOUT_SECONDS}s"
        return probe
    probe.pages = box.get("pages")  # type: ignore[assignment]
    probe.first_text = str(box.get("text", ""))
    probe.meta_title = str(box.get("title", ""))
    probe.error = str(box.get("error", ""))
    return probe


def guess_title_year(probe: PdfProbe) -> Tuple[str, str]:
    title = probe.meta_title.strip()
    if not title or len(title) < 6 or title.lower().startswith(("microsoft", "untitled")):
        lines = [ln.strip() for ln in probe.first_text.splitlines() if len(ln.strip()) > 8]
        title = lines[0][:160] if lines else ""
    years = [int(y.group(0)) for y in YEAR_RE.finditer(probe.first_text[:4000])]
    years = [y for y in years if 1990 <= y <= datetime.now().year + 1]
    return title, (str(max(years)) if years else "")


# --------------------------------------------------------------------------
# Link checking (optional)
# --------------------------------------------------------------------------

def check_url(url: str, timeout: int = 15) -> str:
    headers = {"User-Agent": "Mozilla/5.0 (Supply Chain Data Review link check)"}
    for method in ("HEAD", "GET"):
        try:
            req = urllib.request.Request(url, headers=headers, method=method)
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return str(resp.status)
        except urllib.error.HTTPError as exc:
            if method == "HEAD" and exc.code in (403, 405, 501):
                continue
            return str(exc.code)
        except Exception as exc:
            if method == "HEAD":
                continue
            return type(exc).__name__
    return "unreachable"


# --------------------------------------------------------------------------
# Publication template Table B1
# --------------------------------------------------------------------------

APA_LIKE_RE = re.compile(r"\((?:(?:19|20)\d{2}[a-z]?|n\.d\.|accessed[^)]*)[^)]*\)", re.I)


def read_publication_b1(template: Path, sheet: str = "Table B1") -> List[Tuple[str, str]]:
    """Return (cell, text) for every APA-looking entry on the template's B1 tab."""
    wb = load_workbook(template, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return []
        ws = wb[sheet]
        out = []
        for row in ws.iter_rows(min_row=10):
            for cell in row:
                v = cell.value
                if isinstance(v, str) and APA_LIKE_RE.search(v):
                    out.append((cell.coordinate, v.strip()))
        return out
    finally:
        wb.close()


def reconcile_publication_b1(register: pd.DataFrame, entries: List[Tuple[str, str]]) -> pd.DataFrame:
    pub_norm = [normalize_apa(t) for _, t in entries]
    used = set()
    rows = []
    for _, r in register.iterrows():
        apa = normalize_apa(r.get("apa_reference", ""))
        if not apa:
            continue
        match_idx, kind, ratio = None, "NOT_IN_TEMPLATE", 0.0
        if apa in pub_norm:
            match_idx, kind, ratio = pub_norm.index(apa), "EXACT", 1.0
        else:
            best = max(
                ((i, SequenceMatcher(None, apa, p).ratio()) for i, p in enumerate(pub_norm)),
                key=lambda x: x[1],
                default=(None, 0.0),
            )
            if best[0] is not None and best[1] >= 0.9:
                match_idx, kind, ratio = best[0], "NEAR_MATCH", round(best[1], 3)
        if match_idx is not None:
            used.add(match_idx)
        rows.append({
            "doc_id": r["doc_id"],
            "register_apa_reference": r.get("apa_reference", ""),
            "template_cell": entries[match_idx][0] if match_idx is not None else "",
            "template_apa_reference": entries[match_idx][1] if match_idx is not None else "",
            "match": kind,
            "similarity": ratio,
        })
    # Second pass: same author, substantially revised text in the template.
    for row in rows:
        if row["match"] != "NOT_IN_TEMPLATE":
            continue
        author = fold(parse_apa(row["register_apa_reference"])["author"])
        best = None
        for i, p in enumerate(pub_norm):
            if i in used or not author or not fold(entries[i][1]).startswith(author):
                continue
            ratio = SequenceMatcher(None, normalize_apa(row["register_apa_reference"]), p).ratio()
            if ratio >= 0.35 and (best is None or ratio > best[1]):
                best = (i, ratio)
        if best:
            used.add(best[0])
            row.update(template_cell=entries[best[0]][0], template_apa_reference=entries[best[0]][1],
                       match="REVISED_IN_TEMPLATE", similarity=round(best[1], 3))
    for i, (cell, text) in enumerate(entries):
        if i not in used:
            rows.append({
                "doc_id": "", "register_apa_reference": "", "template_cell": cell,
                "template_apa_reference": text, "match": "TEMPLATE_ONLY", "similarity": 0.0,
            })
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------

def validate_register(
    input_data: Dict[str, List[Dict[str, str]]],
    pdf_folder: Path,
    work_dir: Path,
    publication_template: Optional[Path] = None,
    check_links: bool = False,
    block_on_missing_pdf: bool = True,
    log: Log = print,
) -> RegisterResult:
    b1 = [r for r in input_data.get("B1_Corpus_Documents", []) if any(v for k, v in r.items() if k != "__excel_row")]
    a1 = input_data.get("A1_Organizations", [])
    exclusions = input_data.get("Exclusions_Duplicates", [])
    staged = {r.get("file_name", "").strip().lower() for r in input_data.get("New_Documents", [])}

    issues: List[RegisterIssue] = []

    def add(sev: str, r: Dict[str, str], check: str, msg: str) -> None:
        issues.append(RegisterIssue(sev, r.get("doc_id", ""), r.get("file_name", ""), check, msg))

    org_ids = {r.get("org_id", "").strip() for r in a1 if r.get("org_id", "").strip()}
    excluded = {
        r.get("file_name", "").strip().lower(): r
        for r in exclusions
        if r.get("file_name", "").strip() and r.get("rule", "").strip().upper().startswith("EXCLUDE")
    }

    log(f"Scanning PDF folder: {pdf_folder}")
    pdfs = discover_pdfs(pdf_folder, work_dir / "_unzipped")
    log(f"Found {sum(len(v) for v in pdfs.values())} PDF file(s); probing registered files...")

    doc_counts: Dict[str, int] = {}
    file_counts: Dict[str, int] = {}
    for r in b1:
        doc_counts[r.get("doc_id", "").strip()] = doc_counts.get(r.get("doc_id", "").strip(), 0) + 1
        fn = r.get("file_name", "").strip().lower()
        if fn:
            file_counts[fn] = file_counts.get(fn, 0) + 1

    probes: Dict[str, PdfProbe] = {}
    to_probe = [paths[0] for name, paths in pdfs.items()]
    with ThreadPoolExecutor(max_workers=4) as pool:
        for probe in pool.map(probe_pdf, to_probe):
            probes[probe.path.name.lower()] = probe

    link_status: Dict[str, str] = {}
    if check_links:
        urls = sorted({r.get("source_url", "").strip() for r in b1 if r.get("source_url", "").strip().lower().startswith("http")})
        log(f"Checking {len(urls)} source URL(s)...")
        with ThreadPoolExecutor(max_workers=8) as pool:
            for url, status in zip(urls, pool.map(check_url, urls)):
                link_status[url] = status

    register_rows = []
    sha_owner: Dict[str, str] = {}
    for r in b1:
        doc_id = r.get("doc_id", "").strip()
        fn = r.get("file_name", "").strip()
        title = r.get("title", "").strip()
        apa = r.get("apa_reference", "").strip()
        url = r.get("source_url", "").strip()
        before = len(issues)

        if not doc_id:
            add("ERROR", r, "doc_id", "Register row has no doc_id.")
        elif doc_counts.get(doc_id, 0) > 1:
            add("ERROR", r, "doc_id", f"doc_id {doc_id} appears {doc_counts[doc_id]} times.")

        org = r.get("publishing_org_id", "").strip()
        if not org:
            add("ERROR", r, "publishing_org_id", "No publishing_org_id.")
        else:
            unknown = [o for o in re.split(r"[;,]", org) if o.strip() and o.strip() not in org_ids]
            if unknown:
                add("ERROR", r, "publishing_org_id", f"publishing_org_id {', '.join(o.strip() for o in unknown)} not in A1_Organizations.")

        probe: Optional[PdfProbe] = None
        pdf_path = ""
        if not fn:
            add("INFO", r, "file_name", "No file_name: listed in Table B1 but not counted (reference only).")
        else:
            if file_counts.get(fn.lower(), 0) > 1:
                add("ERROR", r, "file_name", f"file_name registered {file_counts[fn.lower()]} times.")
            if fn.lower() in excluded:
                rule = excluded[fn.lower()]
                keep = rule.get("retain_file_name", "").strip()
                add("ERROR", r, "exclusions", f"File is listed in Exclusions_Duplicates ({rule.get('rule', '')})"
                    + (f"; retain {keep} instead" if keep else "") + ".")
            paths = pdfs.get(fn.lower(), [])
            if not paths:
                add("ERROR" if block_on_missing_pdf else "WARNING", r, "pdf", "PDF not found in the corpus folder.")
            else:
                if len(paths) > 1:
                    add("WARNING", r, "pdf", f"{len(paths)} files share this name; using {paths[0]}.")
                pdf_path = str(paths[0])
                probe = probes.get(fn.lower())
                if probe and probe.error:
                    add("ERROR", r, "pdf_read", f"PDF could not be read: {probe.error}")
                elif probe:
                    if probe.sha256 in sha_owner:
                        add("WARNING", r, "duplicate_content", f"Same file content as {sha_owner[probe.sha256]}.")
                    else:
                        sha_owner[probe.sha256] = doc_id
                    if not probe.first_text.strip():
                        add("WARNING", r, "text_layer", "No text on the first pages (scanned PDF?); OCR before counting.")

        apa_parts = parse_apa(apa)
        if not apa:
            add("WARNING", r, "apa_reference", "No APA reference.")
        else:
            if not apa_parts["author"]:
                add("WARNING", r, "apa_format", "APA reference does not follow 'Author. (Date). Title.' form.")
            y = clean_year(r.get("year", ""))
            if apa_parts["year"] and y and apa_parts["year"] != y:
                add("WARNING", r, "apa_year", f"APA year {apa_parts['year']} differs from year column {y}.")
            if title:
                t_in_apa = fold(title) in fold(apa) or SequenceMatcher(None, fold(title), fold(apa_parts["title"])).ratio() >= 0.85
                if not t_in_apa:
                    add("WARNING", r, "apa_title", "APA reference does not contain the title column.")
            if url and apa_parts["url"] and normalize_url(url) != normalize_url(apa_parts["url"]):
                add("WARNING", r, "apa_url", "URL in APA reference differs from source_url.")
            if url and not apa_parts["url"]:
                add("INFO", r, "apa_url", "APA reference has no URL; source_url available.")
        if url and not re.match(r"^https?://\S+$", url, re.I):
            add("WARNING", r, "source_url", "source_url is not a well-formed http(s) URL.")
        if "verify" in r.get("apa_completion_status", "").lower():
            add("INFO", r, "apa_completion_status", r.get("apa_completion_status", ""))
        if url in link_status and not link_status[url].startswith(("2", "3")):
            add("WARNING", r, "link", f"source_url check returned {link_status[url]}.")

        if title and fn and fold(title).replace(" ", "") == fold(Path(fn).stem).replace(" ", ""):
            add("WARNING", r, "title_placeholder", "Title is just the file name; replace with the document's real title.")
        score: Optional[float] = None
        if probe and not probe.error and title:
            score = title_match_score(title, probe.first_text + "\n" + probe.meta_title)
            if score is not None and score < TITLE_MATCH_MIN and probe.first_text.strip():
                add("WARNING", r, "title_on_pdf",
                    f"Only {int(score * 100)}% of title words found on the PDF's first {PROBE_PAGES} pages; "
                    "check the file/metadata pairing.")

        row_issues = issues[before:]
        sev = {i.severity for i in row_issues}
        status = "ERROR" if "ERROR" in sev else "CHECK" if "WARNING" in sev else ("REFERENCE_ONLY" if not fn else "OK")
        register_rows.append({
            "doc_id": doc_id,
            "register_status": status,
            "file_name": fn,
            "pdf_path": pdf_path,
            "title": title,
            "year": clean_year(r.get("year", "")),
            "authors_or_orgs": r.get("authors_or_orgs", ""),
            "publishing_org_id": org,
            "source_url": url,
            "apa_reference": apa,
            "in_text_citation": r.get("in_text_citation", ""),
            "title_match_on_pdf": "" if score is None else score,
            "pdf_pages": probe.pages if probe else "",
            "pdf_sha256": probe.sha256 if probe else "",
            "link_status": link_status.get(url, ""),
            "issues": "; ".join(f"[{i.severity}] {i.check}: {i.message}" for i in row_issues),
            "input_excel_row": r.get("__excel_row", ""),
        })
    register = pd.DataFrame(register_rows)

    # Unregistered PDFs.
    registered = {r.get("file_name", "").strip().lower() for r in b1}
    unreg_rows = []
    for name, paths in sorted(pdfs.items()):
        if name in registered:
            continue
        probe = probes.get(name)
        if name in excluded:
            rule = excluded[name]
            state = f"excluded ({rule.get('rule', '')})"
        elif probe and probe.sha256 in sha_owner:
            state = f"duplicate content of {sha_owner[probe.sha256]}"
        elif name in staged:
            state = "staged in New_Documents; add to B1"
        else:
            state = "NOT REGISTERED - add to B1 or Exclusions_Duplicates"
        guess_title, guess_year = guess_title_year(probe) if probe else ("", "")
        unreg_rows.append({
            "file_name": paths[0].name,
            "status": state,
            "pdf_path": str(paths[0]),
            "pdf_pages": probe.pages if probe else "",
            "pdf_sha256": probe.sha256 if probe else "",
            "guessed_title": guess_title,
            "guessed_year": guess_year,
        })
        if state.startswith("NOT REGISTERED") or state.startswith("staged"):
            issues.append(RegisterIssue("WARNING", "", paths[0].name, "unregistered_pdf",
                                        f"PDF in corpus folder is not in the register ({state}); it will not be counted."))
    unregistered = pd.DataFrame(unreg_rows, columns=["file_name", "status", "pdf_path", "pdf_pages", "pdf_sha256", "guessed_title", "guessed_year"])

    recon = pd.DataFrame()
    if publication_template and publication_template.exists():
        entries = read_publication_b1(publication_template)
        if entries:
            recon = reconcile_publication_b1(register, entries)
            for _, row in recon.iterrows():
                if row["match"] == "NOT_IN_TEMPLATE":
                    issues.append(RegisterIssue("INFO", row["doc_id"], "", "publication_b1",
                                                "APA reference is not in the publication template's Table B1."))
                elif row["match"] == "TEMPLATE_ONLY":
                    issues.append(RegisterIssue("INFO", "", "", "publication_b1",
                                                f"Template Table B1 {row['template_cell']} has no register row: {row['template_apa_reference'][:90]}"))
            # Proposed APA text: adopt the template wording where it revised the same reference.
            revised = recon[recon["match"].isin(["NEAR_MATCH", "REVISED_IN_TEMPLATE"])]
            proposal = dict(zip(revised["doc_id"], revised["template_apa_reference"]))
            register["proposed_apa_reference"] = register["doc_id"].map(proposal).fillna("")
            for _, row in revised.iterrows():
                issues.append(RegisterIssue("INFO", row["doc_id"], "", "publication_b1",
                                            f"Template Table B1 has revised wording ({row['match']}); see proposed_apa_reference."))
            log(f"Publication Table B1: {len(entries)} entries; "
                + ", ".join(f"{k}={v}" for k, v in recon["match"].value_counts().items()))
        else:
            log("Publication template has no APA entries on 'Table B1'; reconciliation skipped.")

    included = register[(register["register_status"].isin(["OK", "CHECK"])) & register["file_name"].ne("") & register["pdf_path"].ne("")]
    manifest = included[["doc_id", "file_name", "pdf_path", "title", "year", "authors_or_orgs",
                         "publishing_org_id", "source_url", "apa_reference", "pdf_sha256"]].copy()
    manifest.insert(1, "previous_doc_id", manifest["doc_id"])
    manifest["status"] = "INCLUDE"

    n = {s: int((register["register_status"] == s).sum()) for s in STATUS_ORDER} if len(register) else {}
    log(f"Register: {len(register)} rows - " + ", ".join(f"{k} {v}" for k, v in n.items())
        + f"; {len(unregistered)} PDF(s) in folder not registered.")
    return RegisterResult(register, manifest, unregistered, issues, recon)


# --------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------

def apa_sort_key(apa: str) -> str:
    return fold(re.sub(r"^[^\w]+", "", str(apa)))


def _write_table(ws, df: pd.DataFrame, link_cols: Dict[str, str] | None = None, widths: Dict[str, int] | None = None) -> None:
    """Write ``df`` with a bold header; ``link_cols`` maps column -> target column (or 'self')."""
    link_cols = link_cols or {}
    widths = widths or {}
    cols = list(df.columns)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for values in df.itertuples(index=False):
        ws.append(["" if (isinstance(v, float) and pd.isna(v)) else v for v in values])
    for col, target in link_cols.items():
        if col not in cols:
            continue
        ci = cols.index(col) + 1
        ti = ci if target == "self" else cols.index(target) + 1
        for row in range(2, ws.max_row + 1):
            dest = ws.cell(row=row, column=ti).value
            if not dest:
                continue
            dest = str(dest)
            if not re.match(r"^(https?|file):", dest, re.I):
                p = Path(dest)
                if not p.is_absolute():
                    continue
                dest = p.as_uri()
            cell = ws.cell(row=row, column=ci)
            cell.hyperlink = dest
            cell.style = "Hyperlink"
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, min(max(len(col) + 2, 12), 40))
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def publication_b1(register: pd.DataFrame) -> pd.DataFrame:
    rows = register[register["register_status"].ne("ERROR") & register["apa_reference"].ne("")]
    rows = rows.assign(_k=rows["apa_reference"].map(apa_sort_key), _y=rows["year"]).sort_values(["_k", "_y"])
    return rows[["apa_reference", "source_url", "doc_id", "file_name"]].reset_index(drop=True)


def write_register_outputs(result: RegisterResult, out_dir: Path, run_label: str = "") -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    reg = result.register.copy()
    if len(reg):
        reg["_o"] = reg["register_status"].map(STATUS_ORDER)
        reg = reg.sort_values(["_o", "doc_id"]).drop(columns="_o")
    issues = pd.DataFrame([asdict(i) for i in result.issues], columns=["severity", "doc_id", "file_name", "check", "message"])
    pub = publication_b1(result.register) if len(result.register) else pd.DataFrame()

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    counts = reg["register_status"].value_counts().to_dict() if len(reg) else {}
    summary = [
        ("Corpus register validation", run_label or datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Register rows", len(reg)),
        ("OK", counts.get("OK", 0)),
        ("CHECK (warnings)", counts.get("CHECK", 0)),
        ("ERROR (blocks a full run)", counts.get("ERROR", 0)),
        ("Reference only (no file)", counts.get("REFERENCE_ONLY", 0)),
        ("Documents that will be counted", len(result.manifest)),
        ("PDFs in folder not in register", len(result.unregistered)),
        ("", ""),
        ("Sheets", "Register = every B1 row with status, hyperlinks and issues; Table B1 = APA list for "
                   "publication (sorted, hyperlinked); Unregistered PDFs = files to add or exclude; "
                   "Issues = one row per finding; Publication B1 check = register vs. template Table B1."),
    ]
    for k, v in summary:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    for c in ws["A"]:
        c.font = Font(bold=True)

    ws = wb.create_sheet("Register")
    _write_table(ws, reg.drop(columns=["pdf_path"]).assign(open_pdf=reg["pdf_path"].map(lambda p: "open" if p else "")) if len(reg) else reg,
                 widths={"title": 50, "apa_reference": 80, "issues": 90, "source_url": 40, "file_name": 40})
    if len(reg):
        cols = [c.value for c in ws[1]]
        st, fn, su, op = (cols.index(x) + 1 for x in ("register_status", "file_name", "source_url", "open_pdf"))
        for i, path in enumerate(reg["pdf_path"].tolist(), start=2):
            ws.cell(row=i, column=st).fill = STATUS_FILL.get(ws.cell(row=i, column=st).value, PatternFill())
            if path:
                ws.cell(row=i, column=op).hyperlink = Path(path).as_uri()
                ws.cell(row=i, column=op).style = "Hyperlink"
            url = ws.cell(row=i, column=su).value
            if url and str(url).lower().startswith("http"):
                ws.cell(row=i, column=su).hyperlink = str(url)
                ws.cell(row=i, column=su).style = "Hyperlink"

    ws = wb.create_sheet("Table B1")
    ws.append(["Table B-1: sectoral guidance documents reviewed"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    for _, r in pub.iterrows():
        ws.append([r["apa_reference"], r["doc_id"], r["file_name"]])
        cell = ws.cell(row=ws.max_row, column=1)
        if str(r["source_url"]).lower().startswith("http"):
            cell.hyperlink = str(r["source_url"])
            cell.style = "Hyperlink"
    ws.column_dimensions["A"].width = 120
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 50

    ws = wb.create_sheet("Unregistered PDFs")
    _write_table(ws, result.unregistered, link_cols={"file_name": "pdf_path"},
                 widths={"file_name": 50, "status": 45, "guessed_title": 60, "pdf_path": 60})

    ws = wb.create_sheet("Issues")
    _write_table(ws, issues, widths={"message": 110, "file_name": 40})

    if len(result.publication_recon):
        ws = wb.create_sheet("Publication B1 check")
        _write_table(ws, result.publication_recon,
                     widths={"register_apa_reference": 70, "template_apa_reference": 70})

    xlsx = out_dir / "Table_B1_Corpus_Register.xlsx"
    wb.save(xlsx)
    reg.to_csv(out_dir / "corpus_register.csv", index=False, encoding="utf-8-sig")
    issues.to_csv(out_dir / "corpus_register_issues.csv", index=False)
    result.unregistered.to_csv(out_dir / "unregistered_pdfs.csv", index=False)
    result.manifest.to_csv(out_dir / "corpus_manifest.csv", index=False)
    if len(result.publication_recon):
        result.publication_recon.to_csv(out_dir / "publication_b1_reconciliation.csv", index=False)
    return xlsx


def write_proposed_tables(b1_rows: List[Dict[str, str]], result: RegisterResult, out_dir: Path) -> List[Path]:
    """Write drop-in replacements for the reference tables touched by validation.

    - B1_corpus_register_proposed.csv: the register with ``apa_reference``
      replaced by the publication template's revised wording where one was
      found (original kept in ``apa_reference_previous``).
    - new_documents_proposed.csv: one stub row per unregistered PDF, with
      guessed title/year, to complete and move into the register.
    Nothing is applied automatically; review, then copy into the tables folder.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    written: List[Path] = []
    proposal: Dict[str, str] = {}
    if "proposed_apa_reference" in result.register.columns:
        proposal = {d: a for d, a in zip(result.register["doc_id"], result.register["proposed_apa_reference"]) if a}
    if b1_rows:
        cols = [c for c in b1_rows[0].keys() if c != "__excel_row"]
        if proposal and "apa_reference_previous" not in cols:
            cols.append("apa_reference_previous")
        rows = []
        for r in b1_rows:
            rec = {c: r.get(c, "") for c in cols}
            new = proposal.get(r.get("doc_id", ""))
            if new and new != r.get("apa_reference", ""):
                rec["apa_reference_previous"] = r.get("apa_reference", "")
                rec["apa_reference"] = new
            rows.append(rec)
        p = out_dir / "B1_corpus_register_proposed.csv"
        pd.DataFrame(rows, columns=cols).to_csv(p, index=False, encoding="utf-8-sig")
        written.append(p)
    todo = result.unregistered[result.unregistered["status"].str.startswith(("NOT REGISTERED", "staged"))]
    stub = pd.DataFrame({
        "file_name": todo["file_name"],
        "title": todo["guessed_title"],
        "year": todo["guessed_year"],
        "authors_or_orgs": "",
        "publishing_org_id": "",
        "source_url": "",
        "apa_reference": "",
        "status": "to_register",
        "notes": "Stub from register validation: confirm title/year, add org, URL and APA, assign a doc_id in B1.",
    })
    p = out_dir / "new_documents_proposed.csv"
    stub.to_csv(p, index=False, encoding="utf-8-sig")
    written.append(p)
    return written
