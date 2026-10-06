"""Fill a publication (figure/table) template with this run's results.

The template is presentation only: captions, source notes and copy-edit
instructions above each table, an "INSERT FIGURE HERE" anchor on figure tabs,
and a header row for each table. The runner ships a generic template; users
may supply their own (for example a publisher's copy-edit workbook) as long as
tabs are named "Table A1", "Table B1", "Table D1", "Figure 1".."Figure 3" and
tables have recognisable header rows.

Rules
-----
- The user's template is never modified; a filled copy is written.
- A tab is refreshed only when this run produced the data for it. Every other
  tab (curated tables, other figures) is left exactly as in the template.
- Before Table D1 is overwritten, its existing values are saved and compared
  with the new counts (``D1_vs_template.csv``).
- A "Run Info" tab records which tabs were refreshed and from which inputs.
"""
from __future__ import annotations

import re
import unicodedata
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font

Log = Callable[[str], None]

FIGURE_ANCHOR = "insert figure here"
HEADER_SCAN_ROWS = 60


def norm(text) -> str:
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


@dataclass
class TableSpec:
    sheet: str
    # template header label (normalised match, first that hits) -> data column
    columns: List[Tuple[List[str], str]]
    link_columns: Dict[str, str] = field(default_factory=dict)  # data column -> URL column
    group_column: Optional[str] = None  # write only on the first row of each group


D1_SPEC = TableSpec(
    "Table D1",
    [
        (["category"], "category"),
        (["term"], "term"),
        (["variants included", "variants"], "variants_included"),
        (["number of reports referencing term", "reports referencing"], "reports_referencing"),
        (["total occurrences of term across corpus", "total occurrences"], "total_occurrences"),
        (["usage rank in category", "rank"], "rank_in_category"),
    ],
    group_column="category",
)
A1_SPEC = TableSpec(
    "Table A1",
    [(["organization", "organisation"], "organization"), (["website url", "website"], "website_url")],
    link_columns={"website_url": "website_url"},
)
B1_HEADERS = ["sectoral guidance documents", "apa reference", "reference"]
FIGURE_TABS = {
    "Figure 1": "DCF_PRISMA_S_Figure_1_jurisdictional_terms.png",
    "Figure 2": "DCF_PRISMA_S_Figure_2_supply_chain_terms.png",
    "Figure 3": "DCF_PRISMA_S_Figure_3_farm_level_terms.png",
}


def find_header(ws, labels: List[List[str]]) -> Optional[Tuple[int, Dict[int, int]]]:
    """Find the row holding all ``labels``; return (row, {label_index: column})."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, HEADER_SCAN_ROWS)):
        cells = {norm(c.value): c.column for c in row if c.value is not None and str(c.value).strip()}
        hit: Dict[int, int] = {}
        for i, options in enumerate(labels):
            for opt in options:
                if opt in cells:
                    hit[i] = cells[opt]
                    break
        if len(hit) == len(labels):
            return row[0].row, hit
    return None


def find_cell(ws, text: str):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, HEADER_SCAN_ROWS)):
        for c in row:
            if c.value is not None and norm(c.value).startswith(norm(text)):
                return c
    return None


def _unmerge_below(ws, first_row: int, cols: List[int]) -> None:
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row >= first_row and any(rng.min_col <= c <= rng.max_col for c in cols):
            ws.unmerge_cells(str(rng))


def _clear_and_write(ws, first_row: int, col_map: Dict[str, int], df: pd.DataFrame,
                     links: Dict[str, str], group_column: Optional[str]) -> int:
    cols = list(col_map.values())
    _unmerge_below(ws, first_row, cols)
    template_style = {c: copy(ws.cell(row=first_row, column=c)._style) for c in cols}
    for r in range(first_row, max(ws.max_row, first_row) + 1):
        for c in cols:
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.hyperlink = None
    previous_group = None
    for i, rec in enumerate(df.to_dict("records")):
        r = first_row + i
        for key, c in col_map.items():
            cell = ws.cell(row=r, column=c)
            cell._style = copy(template_style[c])
            value = rec.get(key, "")
            if group_column and key == group_column:
                value = value if value != previous_group else None
            cell.value = None if (value is None or (isinstance(value, float) and pd.isna(value)) or value == "") else value
            url = rec.get(links.get(key, ""), "") if key in links else ""
            if url and str(url).lower().startswith("http"):
                cell.hyperlink = str(url)
                cell.font = Font(color="0563C1", underline="single")
        if group_column:
            previous_group = rec.get(group_column)
    return len(df)


def read_existing_d1(ws) -> pd.DataFrame:
    found = find_header(ws, [o for o, _ in D1_SPEC.columns])
    if not found:
        return pd.DataFrame()
    hdr, hit = found
    keys = [k for _, k in D1_SPEC.columns]
    rows, category = [], ""
    for r in range(hdr + 1, ws.max_row + 1):
        vals = {keys[i]: ws.cell(row=r, column=c).value for i, c in hit.items()}
        if not vals.get("term"):
            continue
        category = vals.get("category") or category
        vals["category"] = category
        rows.append(vals)
    return pd.DataFrame(rows)


def compare_d1(old: pd.DataFrame, new: pd.DataFrame) -> pd.DataFrame:
    if old.empty or new.empty:
        return pd.DataFrame()
    o = old.assign(_k=old["term"].map(norm))[["_k", "category", "term", "reports_referencing", "total_occurrences"]]
    n = new.assign(_k=new["term"].map(norm))[["_k", "category", "term", "reports_referencing", "total_occurrences"]]
    m = o.merge(n, on="_k", how="outer", suffixes=("_template", "_run"))
    for col in ["reports_referencing", "total_occurrences"]:
        m[f"{col}_delta"] = pd.to_numeric(m[f"{col}_run"], errors="coerce") - pd.to_numeric(m[f"{col}_template"], errors="coerce")
    m["term"] = m["term_run"].fillna(m["term_template"])
    m["status"] = m.apply(lambda r: "template_only" if pd.isna(r["term_run"]) else "new_in_run" if pd.isna(r["term_template"]) else "both", axis=1)
    cols = ["term", "status", "category_template", "category_run", "reports_referencing_template", "reports_referencing_run",
            "reports_referencing_delta", "total_occurrences_template", "total_occurrences_run", "total_occurrences_delta"]
    return m[cols].sort_values(["status", "term"])


def build_publication_pack(
    template: Path,
    out_xlsx: Path,
    d1: Optional[pd.DataFrame],
    a1: Optional[pd.DataFrame],
    b1: Optional[pd.DataFrame],
    figure_dir: Optional[Path],
    run_info: Dict[str, str],
    log: Log = print,
) -> Dict[str, object]:
    """Fill a copy of ``template``; return a summary dict."""
    wb = load_workbook(template)
    refreshed: List[str] = []
    skipped: List[str] = []
    d1_compare = pd.DataFrame()

    if d1 is not None and not d1.empty and D1_SPEC.sheet in wb.sheetnames:
        ws = wb[D1_SPEC.sheet]
        found = find_header(ws, [o for o, _ in D1_SPEC.columns])
        if found:
            d1_compare = compare_d1(read_existing_d1(ws), d1)
            hdr, hit = found
            col_map = {D1_SPEC.columns[i][1]: c for i, c in hit.items()}
            n = _clear_and_write(ws, hdr + 1, col_map, d1, {}, D1_SPEC.group_column)
            refreshed.append(f"{D1_SPEC.sheet} ({n} terms)")
        else:
            skipped.append(f"{D1_SPEC.sheet}: header row not recognised")

    if a1 is not None and not a1.empty and A1_SPEC.sheet in wb.sheetnames:
        ws = wb[A1_SPEC.sheet]
        found = find_header(ws, [o for o, _ in A1_SPEC.columns])
        if found:
            hdr, hit = found
            col_map = {A1_SPEC.columns[i][1]: c for i, c in hit.items()}
            n = _clear_and_write(ws, hdr + 1, col_map, a1, A1_SPEC.link_columns, None)
            refreshed.append(f"{A1_SPEC.sheet} ({n} organizations)")
        else:
            skipped.append(f"{A1_SPEC.sheet}: header row not recognised")

    if b1 is not None and not b1.empty and "Table B1" in wb.sheetnames:
        ws = wb["Table B1"]
        anchor = None
        for label in B1_HEADERS:
            anchor = find_cell(ws, label)
            if anchor is not None:
                break
        if anchor is not None:
            n = _clear_and_write(ws, anchor.row + 1, {"apa_reference": anchor.column}, b1,
                                 {"apa_reference": "source_url"}, None)
            refreshed.append(f"Table B1 ({n} references)")
        else:
            skipped.append("Table B1: no 'Sectoral Guidance Documents' / 'APA reference' heading")

    if figure_dir is not None:
        for tab, png in FIGURE_TABS.items():
            img_path = figure_dir / png
            if tab not in wb.sheetnames or not img_path.exists():
                continue
            ws = wb[tab]
            anchor = find_cell(ws, FIGURE_ANCHOR)
            row = anchor.row + 1 if anchor is not None else ws.max_row + 2
            col = anchor.column_letter if anchor is not None else "A"
            ws._images = []  # replace any figure placed by an earlier fill
            img = XLImage(str(img_path))
            scale = min(1.0, 640 / max(img.width, 1))
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            ws.add_image(img, f"{col}{row}")
            refreshed.append(f"{tab} (image)")

    untouched = [s for s in wb.sheetnames if not any(r.startswith(s + " ") for r in refreshed)]
    if "Run Info" in wb.sheetnames:
        del wb["Run Info"]
    ws = wb.create_sheet("Run Info")
    lines = [("Filled by", "Supply Chain Data Review desktop runner"),
             ("Filled at", datetime.now().strftime("%Y-%m-%d %H:%M")),
             ("Template", str(template))]
    lines += list(run_info.items())
    lines += [("Refreshed tabs", "; ".join(refreshed) or "none"),
              ("Left as in template", "; ".join(u for u in untouched if u != "Run Info"))]
    lines += [("Skipped", s) for s in skipped]
    for k, v in lines:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120
    for c in ws["A"]:
        c.font = Font(bold=True)

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    load_workbook(out_xlsx, read_only=True).close()  # fail loudly if the file is unreadable
    for line in refreshed:
        log(f"  refreshed {line}")
    for line in skipped:
        log(f"  skipped {line}")
    return {"refreshed": refreshed, "skipped": skipped, "d1_compare": d1_compare}
